"""XLSX workbook → canonical blocks.

A workbook declares sheets, cells, merged ranges and stored values; it does not declare
which cells form a table or which row is a header. So the reader keeps the layout and
decides nothing semantic, applying the same structural rules the DOCX reader applies to a
table, and leaves meaning to the models that read the passages:

* a sheet's filled cells split into regions at fully blank rows, and side by side at a
  fully blank column when the two sides cover different rows - a spacer column inside
  one table spans the same rows on both sides and does not split it;
* a region's leading row whose filled cells are one source cell starting at the region's
  first column - a lone value, or one merged range across the table - is its title;
* in a multi-column region the first row supplies the column headers, extended by the
  next row while a header row holds a cell merged across columns (a group label over its
  columns, joined into "Targets › Minimum"), and each later row is one `table_row`
  passage, serialized exactly as a DOCX row is;
* a one-column region gives one paragraph per cell, a one-row region one paragraph.

Every passage carries the sheet and region title as its heading stack and its exact
`cell_range`, so a row is read with what it is about and traced to where it sits. Hidden
sheets and rows are read and marked, never dropped: hidden content is often superseded,
and a reader told so can weigh it. A region is layout found by blank rows and columns, not
a declared table, so it creates no table group.
"""

from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..models import ContentBlock
from . import rasterizer
from .image_assets import image_asset_byte_size, image_asset_from_bytes
from .table_structure import serialize_table_row

MAX_XLSX_BYTES = 20 * 1024 * 1024
#: Filled rows across the workbook. A raw-data sheet past this is refused whole rather
#: than read in part, as an over-limit PDF is: a partial read looks like a complete one.
MAX_XLSX_ROWS = 2_000
#: Cells a sheet's used range may span. Formatting alone can stretch a sheet's dimension
#: to a million rows; walking that would stall the parse before any bound applied.
MAX_SHEET_CELLS = 2_000_000
#: Rows a header band may span. A merged group label over two header rows is common; a
#: band that kept growing would swallow data rows that happen to hold a merged cell.
MAX_HEADER_ROWS = 3
_IMAGE_MEDIA_TYPES = {"png": "image/png", "jpeg": "image/jpeg", "jpg": "image/jpeg", "gif": "image/gif"}


class WorkbookInputError(ValueError):
    """A workbook cannot be read under the XLSX contract."""


def parse_xlsx(file_path: str, doc_id: str) -> list[ContentBlock]:
    path = Path(file_path)
    if path.stat().st_size > MAX_XLSX_BYTES:
        raise WorkbookInputError(f"{doc_id}: workbook exceeds 20 MB. Split it into smaller files.")
    try:
        values = load_workbook(path, data_only=True)
        formulas = load_workbook(path, data_only=False)
    except Exception as exc:  # noqa: BLE001 - any decoder failure rejects the whole input
        raise WorkbookInputError(f"{doc_id}: not a readable XLSX workbook. Upload a fresh export.") from exc

    # Bounded before anything walks the cells: formatting alone can stretch a sheet's
    # dimension to a million rows, and every later step reads cell by cell.
    for sheet in values.worksheets:
        if sheet.max_row * sheet.max_column > MAX_SHEET_CELLS:
            raise WorkbookInputError(
                f"{doc_id}: sheet '{sheet.title}' spans too many cells to read. "
                "Delete empty formatted rows and columns, or split the workbook."
            )
    _require_calculated(values, formulas, doc_id)
    blocks: list[ContentBlock] = []
    filled_rows = 0
    published_image_bytes = 0
    for sheet_index, sheet in enumerate(values.worksheets, 1):
        grid, anchors, spans = _grid(sheet)
        filled_rows += len(grid)
        if filled_rows > MAX_XLSX_ROWS:
            raise WorkbookInputError(
                f"{doc_id}: workbook has more than {MAX_XLSX_ROWS:,} filled rows. "
                "Split it, or keep the sheets the review needs."
            )
        sheet_meta = {"sheet": sheet.title, "sheet_index": sheet_index}
        if sheet.sheet_state != "visible":
            sheet_meta["hidden_sheet"] = True
        hidden_rows = {row for row, dimension in sheet.row_dimensions.items() if dimension.hidden}
        for region in _regions(grid):
            blocks.extend(_region_blocks(
                doc_id, region, anchors, spans, hidden_rows, sheet.title, sheet_meta,
            ))
        for image in getattr(sheet, "_images", []):
            block, size = _image_block(doc_id, image, sheet.title, sheet_meta)
            published_image_bytes += size
            if published_image_bytes > rasterizer.MAX_RENDERED_BYTES:
                raise WorkbookInputError(f"{doc_id}: workbook images exceed the encoded-image limit.")
            blocks.append(block)
        for _chart in getattr(sheet, "_charts", []):
            blocks.append(ContentBlock(
                id="", doc_id=doc_id, ordinal=-1, block_type="image",
                content="[unsupported visual object]", heading_stack=[sheet.title],
                structural_meta={
                    **sheet_meta,
                    "extraction_warnings": ["unsupported_document_visual"],
                    "unsupported_visual_kind": "chart",
                },
                style_hint={"source": "xlsx_unsupported_visual"},
            ))
    if not blocks:
        raise WorkbookInputError(f"{doc_id}: workbook has no filled cells.")
    for ordinal, block in enumerate(blocks):
        block.ordinal = ordinal
        block.id = f"{doc_id}/b-{ordinal:04d}"
    return blocks


def _require_calculated(values: Any, formulas: Any, doc_id: str) -> None:
    """Refuse a workbook whose formulas were never calculated.

    Excel stores each formula's result when it saves; a file written by a library and
    never opened stores none, and reading it would report every computed value as empty.
    One formula without a result is not that signal - `=IF(x, "", ...)` legitimately
    shows nothing - so the workbook is refused only when no formula has one.
    """
    formula_cells = [
        (sheet.title, cell.coordinate)
        for sheet in formulas.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if getattr(cell, "data_type", None) == "f"
    ]
    if formula_cells and all(values[title][coordinate].value is None for title, coordinate in formula_cells):
        raise WorkbookInputError(
            f"{doc_id}: the workbook's formulas have no saved results. "
            "Open it in Excel, save it, and upload again."
        )


Grid = dict[int, dict[int, str]]
Region = list[tuple[int, dict[int, str]]]


def _grid(sheet: Any) -> tuple[Grid, dict[tuple[int, int], tuple[int, int]], dict[tuple[int, int], int]]:
    """Each filled cell's displayed text by row and column, merged ranges repeated; each
    filled cell's source cell (a merged range's cells share one); and the number of columns
    each merged source spans."""
    merged: dict[tuple[int, int], tuple[int, int]] = {}
    spans: dict[tuple[int, int], int] = {}
    for cell_range in sheet.merged_cells.ranges:
        spans[(cell_range.min_row, cell_range.min_col)] = cell_range.max_col - cell_range.min_col + 1
        for row in range(cell_range.min_row, cell_range.max_row + 1):
            for column in range(cell_range.min_col, cell_range.max_col + 1):
                merged[(row, column)] = (cell_range.min_row, cell_range.min_col)
    grid: dict[int, dict[int, str]] = {}
    anchors: dict[tuple[int, int], tuple[int, int]] = {}
    for row in sheet.iter_rows():
        for cell in row:
            if not hasattr(cell, "column"):
                continue
            anchor = merged.get((cell.row, cell.column))
            source = sheet.cell(*anchor) if anchor else cell
            if text := _cell_text(source):
                grid.setdefault(cell.row, {})[cell.column] = text
                anchors[(cell.row, cell.column)] = (source.row, source.column)
    return grid, anchors, spans


def _cell_text(cell: Any) -> str:
    """A cell's value as Excel shows the kinds it can state exactly."""
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, (int, float)):
        if "%" in (cell.number_format or ""):
            return f"{_number(value * 100)}%"
        return _number(value)
    return " ".join(str(value).split())


def _number(value: float) -> str:
    return str(int(value)) if float(value).is_integer() else format(value, ".15g")


def _regions(grid: Grid) -> list[Region]:
    """Tables as the layout separates them: by blank rows, then by blank columns."""
    regions: list[Region] = []
    for band in _row_runs(grid):
        groups = _column_groups(band)
        if len(groups) == 1:
            regions.append(band)
            continue
        for columns in groups:
            # Within one side, a row it leaves empty separates its own tables.
            side = {row: {c: cells[c] for c in cells if c in columns} for row, cells in band}
            regions.extend(_row_runs({row: cells for row, cells in side.items() if cells}))
    return regions


def _row_runs(grid: Grid) -> list[Region]:
    """Runs of consecutive filled rows; a fully blank row ends a run."""
    runs: list[Region] = []
    previous = None
    for row in sorted(grid):
        if previous is None or row != previous + 1:
            runs.append([])
        runs[-1].append((row, grid[row]))
        previous = row
    return runs


def _column_groups(band: Region) -> list[set[int]]:
    """Column ranges separated by a column the band never fills, kept apart only when the
    two sides cover different rows. A spacer column inside one table has the same rows on
    both sides; a side table beside it starts or ends on other rows."""
    filled = sorted({column for _, cells in band for column in cells})
    groups: list[list[int]] = [[filled[0]]]
    for column in filled[1:]:
        if column == groups[-1][-1] + 1:
            groups[-1].append(column)
        else:
            groups.append([column])

    def rows_of(columns: list[int]) -> frozenset[int]:
        return frozenset(row for row, cells in band if any(c in cells for c in columns))

    merged: list[list[int]] = [groups[0]]
    for group in groups[1:]:
        if rows_of(group) == rows_of(merged[-1]):
            merged[-1].extend(group)
        else:
            merged.append(group)
    return [set(group) for group in merged]


def _region_blocks(
    doc_id: str,
    region: Region,
    anchors: dict[tuple[int, int], tuple[int, int]],
    spans: dict[tuple[int, int], int],
    hidden_rows: set[int],
    sheet_title: str,
    sheet_meta: dict[str, Any],
) -> list[ContentBlock]:
    def meta(row: int, **fields: Any) -> dict[str, Any]:
        return {**sheet_meta, **({"hidden_row": True} if row in hidden_rows else {}), **fields}

    def row_sources(row: int, cells: dict[int, str]) -> set[tuple[int, int]]:
        """The source cells of this region's part of a row."""
        return {anchors[(row, column)] for column in cells}

    heading_stack = [sheet_title]
    rows = region
    region_first = min(column for _, cells in rows for column in cells)
    leading = row_sources(*rows[0]) if rows else set()
    if len(rows) > 1 and len(leading) == 1 and min(rows[0][1]) == region_first:
        title_row, cells = rows[0]
        title = next(iter(cells.values()))
        heading_stack = [sheet_title, title]
        rows = rows[1:]
        title_block = _block(
            doc_id, "heading", title, [sheet_title],
            meta(title_row, cell_range=_range(title_row, min(cells), max(cells))), "xlsx_region_title",
        )
    else:
        title_block = None

    first = min(column for _, cells in rows for column in cells)
    last = max(column for _, cells in rows for column in cells)
    blocks: list[ContentBlock] = [title_block] if title_block else []
    if first == last:
        blocks += [
            _block(doc_id, "paragraph", cells[first], heading_stack,
                   meta(row, cell_range=_range(row, first, first)), "xlsx_column")
            for row, cells in rows
        ]
        return blocks
    if len(rows) == 1:
        row, cells = rows[0]
        blocks.append(_block(
            doc_id, "paragraph", " | ".join(cells[column] for column in sorted(cells)), heading_stack,
            meta(row, cell_range=_range(row, first, last)), "xlsx_row",
        ))
        return blocks

    columns = range(first, last + 1)
    # A header row holding a cell merged across columns is a group label over the row
    # beneath it, so that row joins the header and each column's label reads top-down.
    band = 1
    while (band < min(MAX_HEADER_ROWS, len(rows) - 1)
           and any(spans.get(source, 1) > 1 for source in row_sources(*rows[band - 1]))):
        band += 1
    header_rows = rows[:band]
    headers = [
        " › ".join(dict.fromkeys(label for _, cells in header_rows if (label := cells.get(column, ""))))
        for column in columns
    ]
    header_row = header_rows[0][0]
    for row, cells in rows[band:]:
        content, table_cells = serialize_table_row(headers, [cells.get(column, "") for column in columns])
        if not content:
            continue
        blocks.append(_block(
            doc_id, "table_row", content, heading_stack,
            meta(
                row,
                cell_range=_range(row, first, last),
                header_range=f"{_range(header_row, first, first)}:{get_column_letter(last)}{header_rows[-1][0]}",
                column_headers=headers,
                table_cells=table_cells,
            ),
            "table_row",
        ))
    return blocks


def _image_block(doc_id: str, image: Any, sheet_title: str, sheet_meta: dict[str, Any]) -> tuple[ContentBlock, int]:
    media_type = _IMAGE_MEDIA_TYPES.get(str(getattr(image, "format", "")).lower(), "")
    asset = image_asset_from_bytes(image._data(), media_type) if media_type else None
    meta = dict(sheet_meta)
    anchor = getattr(getattr(image, "anchor", None), "_from", None)
    if anchor is not None:
        meta["cell_range"] = f"{get_column_letter(anchor.col + 1)}{anchor.row + 1}"
    if asset is None:
        meta["extraction_warnings"] = ["document_image_unavailable"]
    block = _block(
        doc_id, "image", "[image]" if asset else "[image unavailable]", [sheet_title], meta, "xlsx_image",
    )
    block.image = asset
    return block, image_asset_byte_size(asset) if asset else 0


def _range(row: int, first: int, last: int) -> str:
    start = f"{get_column_letter(first)}{row}"
    return start if first == last else f"{start}:{get_column_letter(last)}{row}"


def _block(
    doc_id: str, block_type: str, content: str, heading_stack: list[str],
    structural_meta: dict[str, Any], source: str,
) -> ContentBlock:
    return ContentBlock(
        id="", doc_id=doc_id, ordinal=-1, block_type=block_type, content=content,
        heading_stack=list(heading_stack), structural_meta=structural_meta,
        style_hint={"source": source},
    )
