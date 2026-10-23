"""A draft as a Word table, for the team to edit into their iTPP.

Every target in it is a suggestion and the document says so; every row names the sources it
rests on, so the table carries its provenance into the document it ends up in.
"""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.shared import Pt

from shared.vocabulary import search_term

from .contract import validate_result_contract

COLUMNS = ("Variable", "Suggested target", "Based on", "Sources")
NOT_SET = "Not set"


def export_docx(result: dict[str, Any]) -> bytes:
    validate_result_contract(result)
    document = Document()
    document.add_heading("Draft iTPP targets", level=1)
    context = f"{result['display_name']} · {_label(search_term(result['indication']))}"
    document.add_paragraph(context)
    document.add_paragraph(
        "Suggested by PDIS Drafter from comparable products' results and published guidance. "
        "Every target is a suggestion to check and edit; each row lists the sources it rests on."
    )
    table = document.add_table(rows=1, cols=len(COLUMNS))
    table.style = "Table Grid"
    for cell, heading in zip(table.rows[0].cells, COLUMNS, strict=True):
        cell.text = heading
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for variable in result["variables"]:
        target, basis, sources = _row(variable)
        cells = table.add_row().cells
        for cell, text in zip(cells, (variable["label"], target, basis, sources), strict=True):
            cell.text = text
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


LEANS_ON_LABEL = {
    "comparable_products": "Comparable products",
    "guidance": "Published guidance",
    "both": "Comparable products and published guidance",
}


def _row(variable: dict[str, Any]) -> tuple[str, str, str]:
    if variable["status"] == "failed":
        return NOT_SET, "Could not be suggested", ""
    suggestion = variable["suggestion"]
    if suggestion["kind"] == "insufficient":
        return NOT_SET, "Not enough evidence", ""
    return _target(suggestion), LEANS_ON_LABEL[suggestion["leans_on"]], _sources(variable, suggestion["basis_ids"])


def _target(values: dict[str, Any]) -> str:
    statement = str(values.get("statement") or "").strip()
    if statement:
        return statement
    parts = [
        f"{label}: {value}"
        for label, value in (("Minimum", values.get("minimum")), ("Preferred", values.get("preferred")))
        if str(value or "").strip()
    ]
    return "\n".join(parts)


def _sources(variable: dict[str, Any], basis_ids: list[str]) -> str:
    by_id = {observation["id"]: observation for observation in variable["observations"]}
    lines = []
    for identifier in basis_ids:
        observation = by_id[identifier]
        lines.append(f"{observation['title']} ({observation['source_label']}) {observation['url']}".strip())
    return "\n".join(dict.fromkeys(lines))


def _label(text: str) -> str:
    return text[:1].upper() + text[1:]
