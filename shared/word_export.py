"""Word's layout for an exported result: one landscape document, tables in small type.

The rows are the tool's (`shared.report_export`); this decides only how Word shows them - a
bold header row in small type on a landscape page, where five columns of prose fit.
"""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Pt

from shared.report_export import CAVEAT, NONE, ExportReport, Table

CELL_POINTS = 9


def word_document(report: ExportReport) -> bytes:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    document.add_heading(report.title, level=1)
    for line in (report.context, report.about, CAVEAT, report.summary, *report.notes):
        if line:
            document.add_paragraph(line)
    for table in report.tables:
        if table.heading:
            document.add_heading(table.heading, level=2)
        if table.note:
            document.add_paragraph(table.note)
        # A table with no rows is its note alone: a header row over nothing reads as lost rows.
        if table.rows:
            _table(document, table, section.page_width - section.left_margin - section.right_margin)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _table(document: Any, table: Table, width: int) -> None:
    grid = document.add_table(rows=1, cols=len(table.columns))
    grid.style = "Table Grid"
    shares = table.widths or (1.0,) * len(table.columns)
    if len(shares) != len(table.columns):
        raise ValueError("a table gives one width per column")
    grid.autofit = False
    for cell, heading in zip(grid.rows[0].cells, table.columns, strict=True):
        cell.text = heading
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(CELL_POINTS)
    for row in table.rows:
        for cell, text in zip(grid.add_row().cells, row, strict=True):
            cell.text = text or NONE
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(CELL_POINTS)
    # Set on the table's column grid and on every cell of the column: Word reads the cells,
    # other readers the grid.
    for index, column in enumerate(grid.columns):
        column_width = int(width * shares[index] / sum(shares))
        column.width = column_width
        for cell in column.cells:
            cell.width = column_width
