from __future__ import annotations

import base64
import re
import tempfile
import unittest
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml

from services.chunker import blocks_to_dicts, run_pipeline
from services.chunker.stages.docx_parts import DocxPartReader, PartWarning
from tests.test_chunker_images import PNG_1X1


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
V = "urn:schemas-microsoft-com:vml"
O = "urn:schemas-microsoft-com:office:office"


class UnsupportedVisualIdentityTests(unittest.TestCase):
    def test_structural_object_types_remain_warned_without_classifying_importance(self):
        cases = [
            (f'<w:object xmlns:w="{W}"/>', "embedded_object"),
            ('<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"/>', "chart"),
            ('<d:relIds xmlns:d="http://schemas.openxmlformats.org/drawingml/2006/diagram"/>', "diagram"),
            (f'<v:line xmlns:v="{V}"/>', "drawing"),
        ]
        for xml, expected in cases:
            with self.subTest(expected=expected):
                items = DocxPartReader(Document()).container_supplements(
                    parse_xml(xml), "/word/header1.xml", "header"
                )
                self.assertEqual(len(items), 1)
                self.assertIsInstance(items[0], PartWarning)
                self.assertEqual(items[0].visual_kind, expected)
                self.assertEqual(items[0].kind, "header")


def _replace_zip_members(path: Path, replacements: dict[str, bytes]) -> None:
    rewritten = path.with_suffix(".rewritten.docx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(rewritten, "w") as target:
        for item in source.infolist():
            target.writestr(item, replacements.get(item.filename, source.read(item.filename)))
        existing = set(source.namelist())
        for name, payload in replacements.items():
            if name not in existing:
                target.writestr(name, payload)
    rewritten.replace(path)


def _coverage_fixture(path: Path) -> None:
    document = Document()
    document.add_paragraph("Body before")
    outer = document.add_table(rows=1, cols=1)
    outer.cell(0, 0).paragraphs[0].add_run("Outer cell")
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "Nested cell"
    ordered = document.add_table(rows=1, cols=1).cell(0, 0)
    ordered.paragraphs[0].text = "Before text box"
    ordered.add_paragraph("Table text-box anchor")
    ordered.add_paragraph("After text box")
    multi = document.add_table(rows=2, cols=2)
    multi.cell(0, 0).text = "H1"
    multi.cell(0, 1).text = "H2"
    multi.cell(1, 0).text = "Left value"
    multi.cell(1, 1).text = "Multi anchor"
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.header.paragraphs[0].text = "Header retained once"
    section.footer.paragraphs[0].text = "Footer retained once"
    document.add_paragraph("Body after")
    document.add_heading("Final section", level=1)
    document.add_picture(BytesIO(PNG_1X1))
    document.save(path)

    with zipfile.ZipFile(path) as package:
        document_xml = package.read("word/document.xml").decode()
        rels_xml = package.read("word/_rels/document.xml.rels").decode()
        content_types = package.read("[Content_Types].xml").decode()
        header_name = next(name for name in package.namelist() if name.startswith("word/header"))
        header_xml = package.read(header_name).decode()
        image_name = next(name for name in package.namelist() if name.startswith("word/media/"))

    textbox = (
        f'<w:p xmlns:w="{W}" xmlns:r="{R}" xmlns:v="{V}" xmlns:o="{O}">'
        '<w:r><w:t>Anchor text</w:t></w:r><w:r><w:pict><v:shape>'
        '<v:textbox><w:txbxContent><w:p><w:r><w:t>Floating box</w:t></w:r>'
        f'</w:p><w:p><w:r><w:drawing><a:blip xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'r:embed="BODY_IMAGE_REL"/></w:drawing></w:r></w:p>'
        '</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>'
    )
    body_image_rel = re.search(
        r'<Relationship Id="([^"]+)"[^>]+Type="[^"]+/image"', rels_xml
    )
    assert body_image_rel is not None
    textbox = textbox.replace("BODY_IMAGE_REL", body_image_rel.group(1))
    note_reference = (
        f'<w:p xmlns:w="{W}"><w:r><w:t>Footnote anchor</w:t></w:r>'
        '<w:r><w:footnoteReference w:id="2"/></w:r></w:p>'
    )
    endnote_reference = (
        f'<w:p xmlns:w="{W}"><w:r><w:t>Endnote anchor</w:t></w:r>'
        '<w:r><w:endnoteReference w:id="3"/></w:r></w:p>'
    )
    repeated_note_reference = (
        f'<w:p xmlns:w="{W}"><w:r><w:footnoteReference w:id="2"/></w:r></w:p>'
    )
    unsupported = (
        f'<w:p xmlns:w="{W}" xmlns:o="{O}"><w:r><w:object>'
        '<o:OLEObject Type="Embed"/></w:object></w:r></w:p>'
    )
    unsupported_native_shapes = (
        f'<w:p xmlns:w="{W}" xmlns:v="{V}" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<w:r><w:pict><v:line/></w:pict></w:r><w:r><w:drawing><wps:wsp>'
        '<wps:spPr><a:prstGeom prst="rect"/></wps:spPr></wps:wsp></w:drawing></w:r></w:p>'
    )
    alternate = (
        f'<w:p xmlns:w="{W}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        '<mc:AlternateContent><mc:Choice Requires="w"><w:r><w:t>Choice branch</w:t></w:r>'
        '</mc:Choice><mc:Fallback><w:r><w:t>Fallback duplicate</w:t></w:r>'
        '</mc:Fallback></mc:AlternateContent></w:p>'
    )
    before_final_section, final_section = document_xml.rsplit("<w:sectPr", 1)
    document_xml = (
        before_final_section
        + textbox
        + note_reference
        + endnote_reference
        + repeated_note_reference
        + unsupported
        + unsupported_native_shapes
        + alternate
        + "<w:sectPr"
        + final_section
    )
    document_xml = document_xml.replace(
        "<w:r><w:t>Outer cell</w:t></w:r>",
        '<w:r><w:t>Outer cell</w:t></w:r><w:r><w:footnoteReference w:id="9"/></w:r>',
        1,
    )
    document_xml = document_xml.replace(
        "<w:r><w:t>Multi anchor</w:t></w:r>",
        f'<w:r><w:t>Multi anchor</w:t></w:r><w:r><w:footnoteReference w:id="10"/>'
        f'</w:r><w:r><w:pict xmlns:v="{V}">'
        '<v:shape><v:textbox><w:txbxContent><w:p><w:r>'
        '<w:t>Multi cell text box</w:t></w:r></w:p></w:txbxContent>'
        '</v:textbox></v:shape><v:rect/></w:pict></w:r>',
        1,
    )
    document_xml = document_xml.replace(
        "<w:r><w:t>Table text-box anchor</w:t></w:r>",
        f'<w:r><w:t>Table text-box anchor</w:t></w:r><w:r><w:pict xmlns:v="{V}">'
        '<v:shape><v:textbox><w:txbxContent><w:p><w:r>'
        '<w:t>Middle text box</w:t></w:r></w:p></w:txbxContent>'
        '</v:textbox></v:shape></w:pict></w:r>',
        1,
    )
    header_textbox = (
        f'<w:p xmlns:w="{W}" xmlns:v="{V}"><w:r><w:pict><v:shape><v:textbox>'
        '<w:txbxContent><w:p><w:r><w:t>Header floating box</w:t></w:r></w:p>'
        '<w:tbl><w:tblPr/><w:tblGrid><w:gridCol/></w:tblGrid><w:tr><w:tc>'
        '<w:tcPr/><w:p><w:r><w:t>Header text-box table</w:t></w:r>'
        '<w:r><w:footnoteReference w:id="11"/></w:r></w:p>'
        '</w:tc></w:tr></w:tbl>'
        '</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>'
    )
    header_xml = header_xml.replace("</w:hdr>", header_textbox + "</w:hdr>")
    rels_xml = rels_xml.replace(
        "</Relationships>",
        '<Relationship Id="rIdFootnotes" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" '
        'Target="footnotes.xml"/><Relationship Id="rIdEndnotes" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes" '
        'Target="endnotes.xml"/></Relationships>',
    )
    content_types = content_types.replace(
        "</Types>",
        '<Override PartName="/word/footnotes.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
        '<Override PartName="/word/endnotes.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml"/>'
        "</Types>",
    )
    footnotes = (
        f'<w:footnotes xmlns:w="{W}"><w:footnote w:id="2"><w:p><w:r>'
        '<w:t>Referenced footnote</w:t></w:r></w:p><w:tbl><w:tblPr/>'
        '<w:tblGrid><w:gridCol/></w:tblGrid><w:tr><w:tc><w:tcPr/><w:p><w:r>'
        '<w:t>Footnote table cell</w:t></w:r><w:r><w:object/></w:r></w:p>'
        '</w:tc></w:tr></w:tbl><w:p><w:r><w:drawing><wp:inline '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
        f'<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:graphicData><a:blip r:embed="rIdNoteImage" xmlns:r="{R}"/>'
        '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p></w:footnote>'
        '<w:footnote w:id="9"><w:p><w:r><w:t>Unreferenced footnote</w:t></w:r>'
        '</w:p></w:footnote><w:footnote w:id="10"><w:p><w:r>'
        '<w:t>Multi cell note</w:t></w:r></w:p></w:footnote></w:footnotes>'
    ).encode()
    footnotes = footnotes.replace(
        b"</w:footnotes>",
        (
            f'<w:footnote w:id="11"><w:p><w:r><w:t>Header table note</w:t>'
            f'</w:r></w:p><w:p><w:r><w:drawing><a:blip xmlns:a="{R.replace("officeDocument/2006/relationships", "drawingml/2006/main")}" '
            f'xmlns:r="{R}" r:embed="rIdNoteImage"/></w:drawing></w:r></w:p>'
            '</w:footnote></w:footnotes>'
        ).encode(),
    )
    endnotes = (
        f'<w:endnotes xmlns:w="{W}"><w:endnote w:id="3"><w:p><w:r>'
        '<w:t>Referenced endnote</w:t></w:r></w:p><w:tbl><w:tblPr/>'
        '<w:tblGrid><w:gridCol/></w:tblGrid><w:tr><w:tc><w:tcPr/><w:p><w:r>'
        '<w:t>Endnote table cell</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
        '</w:endnote></w:endnotes>'
    ).encode()
    footnote_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'<Relationship Id="rIdNoteImage" Type="{R}/image" Target="media/{Path(image_name).name}"/>'
        '</Relationships>'
    ).encode()
    _replace_zip_members(
        path,
        {
            "word/document.xml": document_xml.encode(),
            "word/_rels/document.xml.rels": rels_xml.encode(),
            "[Content_Types].xml": content_types.encode(),
            "word/footnotes.xml": footnotes,
            "word/endnotes.xml": endnotes,
            "word/_rels/footnotes.xml.rels": footnote_rels,
            header_name: header_xml.encode(),
        },
    )


class ChunkerDocxPartCoverageTests(unittest.TestCase):
    def test_retains_explicit_docx_parts_once_with_stable_locations(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "coverage.docx"
            _coverage_fixture(path)
            blocks = run_pipeline(str(path), "coverage")

        contents = [block.content for block in blocks]
        self.assertEqual(contents.count("Floating box"), 1)
        self.assertEqual(contents.count("Referenced footnote"), 1)
        self.assertEqual(contents.count("Anchor text"), 1)
        self.assertEqual(
            [content for content in contents if "Floating box" in content],
            ["Floating box"],
        )
        self.assertEqual(contents.count("Header retained once"), 1)
        self.assertEqual(contents.count("Footer retained once"), 1)
        self.assertEqual(contents.count("Referenced footnote"), 1)
        self.assertEqual(contents.count("Referenced endnote"), 1)
        self.assertIn("Unreferenced footnote", contents)
        self.assertIn("Outer cell", contents)
        self.assertIn("Nested cell", contents)
        self.assertIn("Footnote table cell", contents)
        self.assertIn("Endnote table cell", contents)
        self.assertIn("Header floating box", contents)
        self.assertIn("Header text-box table", contents)
        self.assertLess(contents.index("Before text box"), contents.index("Middle text box"))
        self.assertLess(contents.index("Middle text box"), contents.index("After text box"))
        multi_row = next(
            block for block in blocks
            if block.block_type == "table_row" and "Multi anchor" in block.content
        )
        self.assertLess(contents.index(multi_row.content), contents.index("Multi cell text box"))
        self.assertLess(contents.index(multi_row.content), contents.index("Multi cell note"))
        self.assertEqual(
            [cell["value"] for cell in multi_row.structural_meta["table_cells"]],
            ["Left value", "Multi anchor"],
        )
        self.assertIn("Choice branch", contents)
        self.assertNotIn("Fallback duplicate", contents)

        by_content = {block.content: block for block in blocks}
        expected_parts = {
            "Floating box": ("/word/document.xml", "textbox"),
            "Referenced footnote": ("/word/footnotes.xml", "footnote"),
            "Referenced endnote": ("/word/endnotes.xml", "endnote"),
        }
        for content, (part, kind) in expected_parts.items():
            block = by_content[content]
            self.assertEqual(block.structural_meta["document_part"], part)
            self.assertEqual(block.structural_meta["document_part_kind"], kind)

        header = by_content["Header retained once"]
        footer = by_content["Footer retained once"]
        self.assertRegex(header.structural_meta["document_part"], r"^/word/header\d+\.xml$")
        self.assertRegex(footer.structural_meta["document_part"], r"^/word/footer\d+\.xml$")
        self.assertEqual(header.structural_meta["document_part_kind"], "header")
        self.assertEqual(footer.structural_meta["document_part_kind"], "footer")
        self.assertEqual(header.heading_stack, [])
        self.assertEqual(footer.heading_stack, [])
        self.assertEqual(len({block.id for block in blocks}), len(blocks))
        self.assertNotIn("page", {key for block in blocks for key in block.structural_meta})

        warning_blocks = [
            block for block in blocks
            if "unsupported_document_visual"
            in block.structural_meta.get("extraction_warnings", [])
        ]
        self.assertEqual(len(warning_blocks), 5)
        self.assertEqual(
            {block.structural_meta.get("unsupported_visual_kind") for block in warning_blocks},
            {"drawing", "embedded_object"},
        )
        self.assertEqual(
            {block.structural_meta["document_part"] for block in warning_blocks},
            {"/word/document.xml", "/word/footnotes.xml"},
        )

        note_images = [
            block for block in blocks
            if block.block_type == "image"
            and block.structural_meta.get("document_part_kind") == "footnote"
            and block.image is not None
        ]
        self.assertEqual(
            [block.structural_meta.get("note_id") for block in note_images],
            ["2", "11"],
        )
        header_note_image = next(
            block for block in blocks
            if block.block_type == "image"
            and block.structural_meta.get("note_id") == "11"
        )
        self.assertEqual(
            header_note_image.structural_meta["document_part"],
            "/word/footnotes.xml",
        )
        self.assertEqual(header_note_image.structural_meta["document_part_kind"], "footnote")
        self.assertIsNotNone(header_note_image.image)
        assert header_note_image.image is not None
        self.assertEqual(base64.b64decode(header_note_image.image.data_base64), PNG_1X1)
        textbox_images = [
            block for block in blocks
            if block.block_type == "image"
            and block.structural_meta.get("document_part_kind") == "textbox"
            and block.image is not None
        ]
        self.assertEqual(len(textbox_images), 1)

        serialized = blocks_to_dicts(blocks)
        serialized_note = next(
            block for block in serialized if block["content"] == "Referenced footnote"
        )
        self.assertEqual(serialized_note["structural_meta"]["note_id"], "2")


if __name__ == "__main__":
    unittest.main()
