"""DOCX structure the parser must read through rather than around.

Content controls (`w:sdt`) and custom XML (`w:customXml`) wrap ordinary content at
block, row, cell and run level; they declare metadata about what they wrap, never a
reason to drop it. Headings are declared by outline level, directly or through the
style chain, not only by a style whose name begins "Heading".
"""

from __future__ import annotations

import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from PIL import Image

from services.chunker import run_pipeline
from services.chunker.stages.image_assets import _convert_raster_to_png
from tests.test_chunker_images import PNG_1X1


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _fragment(xml: str):
    return parse_xml(f'<w:root xmlns:w="{W}">{xml}</w:root>')[0]


def _wrap(element, wrapper: str) -> None:
    """Move `element` inside a new `w:sdt` or `w:customXml` at its position."""
    if wrapper == "sdt":
        container = parse_xml(
            f'<w:sdt xmlns:w="{W}"><w:sdtPr><w:alias w:val="Control"/></w:sdtPr>'
            "<w:sdtContent/></w:sdt>"
        )
        inner = container.find(qn("w:sdtContent"))
    else:
        container = parse_xml(
            f'<w:customXml xmlns:w="{W}" w:element="field"><w:customXmlPr/></w:customXml>'
        )
        inner = container
    element.addprevious(container)
    inner.append(element)


def _parse(document) -> list:
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "structure.docx"
        document.save(path)
        return run_pipeline(str(path), "structure")


class TransparentContainerTests(unittest.TestCase):
    def test_block_level_content_controls_and_custom_xml_are_read(self) -> None:
        document = Document()
        document.add_paragraph("Before controls")
        in_control = document.add_paragraph("Inside block control")
        _wrap(in_control._p, "sdt")
        in_custom = document.add_paragraph("Inside custom XML")
        _wrap(in_custom._p, "custom")
        table = document.add_table(rows=2, cols=2)
        for (row, column), value in {
            (0, 0): "Measure", (0, 1): "Target",
            (1, 0): "Efficacy", (1, 1): "75%",
        }.items():
            table.cell(row, column).text = value
        _wrap(table._tbl, "sdt")
        document.add_paragraph("After controls")

        blocks = _parse(document)
        contents = [block.content for block in blocks]

        self.assertEqual(
            [content for content in contents if "Target" not in content],
            ["Before controls", "Inside block control", "Inside custom XML", "After controls"],
        )
        row = next(block for block in blocks if block.block_type == "table_row")
        self.assertEqual(row.content, "Measure: Efficacy, Target: 75%")
        self.assertLess(contents.index("Inside custom XML"), contents.index(row.content))
        self.assertLess(contents.index(row.content), contents.index("After controls"))

    def test_row_and_cell_controls_are_read(self) -> None:
        document = Document()
        table = document.add_table(rows=3, cols=2)
        for (row, column), value in {
            (0, 0): "Measure", (0, 1): "Target",
            (1, 0): "Efficacy", (1, 1): "75%",
            (2, 0): "Doses", (2, 1): "Two",
        }.items():
            table.cell(row, column).text = value
        second_row, third_row = table.rows[1]._tr, table.rows[2]._tr
        _wrap(second_row, "sdt")
        _wrap(third_row.tc_lst[1], "custom")

        rows = [block.content for block in _parse(document) if block.block_type == "table_row"]

        self.assertEqual(
            rows,
            ["Measure: Efficacy, Target: 75%", "Measure: Doses, Target: Two"],
        )

    def test_inline_wrappers_in_cells_keep_their_text(self) -> None:
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Field"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Site"
        paragraph = table.cell(1, 1).paragraphs[0]._p
        for fragment in (
            '<w:sdt><w:sdtPr/><w:sdtContent><w:r><w:t xml:space="preserve">Seattle </w:t>'
            "</w:r></w:sdtContent></w:sdt>",
            '<w:ins w:id="1" w:author="a"><w:r><w:t xml:space="preserve">and </w:t></w:r></w:ins>',
            '<w:smartTag w:uri="u" w:element="place"><w:r><w:t xml:space="preserve">Nairobi </w:t>'
            "</w:r></w:smartTag>",
            '<w:fldSimple w:instr="PAGE"><w:r><w:t>7</w:t></w:r></w:fldSimple>',
        ):
            paragraph.append(_fragment(fragment))
        single = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]._p
        single.append(_fragment(
            '<w:ins w:id="2" w:author="a"><w:r><w:t>Single cell insert</w:t></w:r></w:ins>'
        ))

        blocks = _parse(document)

        row = next(block for block in blocks if block.block_type == "table_row")
        self.assertEqual(row.content, "Field: Site, Value: Seattle and Nairobi 7")
        self.assertIn("Single cell insert", [block.content for block in blocks])


class NestedTableTests(unittest.TestCase):
    def test_nested_table_in_a_multi_column_cell_is_emitted_once(self) -> None:
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Measure"
        table.cell(0, 1).text = "Target"
        table.cell(1, 0).text = "Efficacy"
        cell = table.cell(1, 1)
        cell.text = "See breakdown"
        nested = cell.add_table(rows=2, cols=2)
        nested.cell(0, 0).text = "Age"
        nested.cell(0, 1).text = "Value"
        nested.cell(1, 0).text = "Infants"
        nested.cell(1, 1).text = "80%"
        nested.cell(1, 1).paragraphs[0].add_run().add_picture(BytesIO(PNG_1X1))
        document.add_paragraph("After table")

        blocks = _parse(document)
        contents = [block.content for block in blocks]

        # python-docx closes a cell after a nested table with an empty paragraph, as
        # Word does; the row's text keeps that cell's own paragraphs as before.
        outer = [content for content in contents if content.startswith("Measure: Efficacy")]
        inner = "Age: Infants, Value: 80%"
        self.assertEqual(len(outer), 1)
        self.assertNotIn("Infants", outer[0])
        self.assertEqual(contents.count(inner), 1)
        self.assertLess(contents.index(outer[0]), contents.index(inner))
        self.assertLess(contents.index(inner), contents.index("After table"))
        self.assertEqual(sum(block.block_type == "image" for block in blocks), 1)
        groups = {
            block.structural_meta.get("table_group")
            for block in blocks if block.content in {outer[0], inner}
        }
        self.assertEqual(len(groups), 1)


class HeadingOutlineTests(unittest.TestCase):
    def test_outline_level_resolves_through_paragraph_and_style_chain(self) -> None:
        document = Document()
        custom = document.styles.add_style("Programme Section", WD_STYLE_TYPE.PARAGRAPH)
        custom.base_style = document.styles["Heading 1"]
        document.add_paragraph("Custom section", style="Programme Section")
        direct = document.add_paragraph("Direct outline")
        direct._p.get_or_add_pPr().append(
            _fragment('<w:outlineLvl w:val="1"/>')
        )
        document.add_paragraph("Body under direct outline")
        demoted = document.add_paragraph("Demoted heading", style="Heading 2")
        demoted._p.get_or_add_pPr().append(
            _fragment('<w:outlineLvl w:val="9"/>')
        )
        document.add_heading("Plain heading", level=3)

        blocks = {block.content: block for block in _parse(document)}

        self.assertEqual(blocks["Custom section"].block_type, "heading")
        self.assertEqual(blocks["Custom section"].structural_meta["heading_level"], 1)
        self.assertEqual(blocks["Direct outline"].block_type, "heading")
        self.assertEqual(blocks["Direct outline"].structural_meta["heading_level"], 2)
        self.assertEqual(
            blocks["Body under direct outline"].heading_stack,
            ["Custom section", "Direct outline"],
        )
        self.assertEqual(blocks["Demoted heading"].block_type, "paragraph")
        self.assertEqual(blocks["Plain heading"].structural_meta["heading_level"], 3)


class DocxImageBoundTests(unittest.TestCase):
    def _document_with_images(self, count: int):
        document = Document()
        for index in range(count):
            document.add_paragraph(f"Figure {index}")
            document.add_picture(BytesIO(PNG_1X1))
        return document

    def test_image_over_the_pixel_bound_is_unavailable_not_retained(self) -> None:
        with patch(
            "services.chunker.stages.rasterizer.MAX_RENDER_PIXELS_PER_PAGE", 0
        ):
            blocks = _parse(self._document_with_images(1))

        image = next(block for block in blocks if block.block_type == "image")
        self.assertIsNone(image.image)
        self.assertIn(
            "document_image_unavailable", image.structural_meta["extraction_warnings"]
        )

    def test_retained_image_bytes_are_bounded_per_document(self) -> None:
        with patch(
            "services.chunker.stages.rasterizer.MAX_RENDERED_BYTES", len(PNG_1X1)
        ):
            blocks = _parse(self._document_with_images(2))

        images = [block for block in blocks if block.block_type == "image"]
        self.assertIsNotNone(images[0].image)
        self.assertIsNone(images[1].image)
        self.assertIn(
            "document_image_unavailable", images[1].structural_meta["extraction_warnings"]
        )

    def test_raster_normalization_refuses_to_decode_past_the_pixel_bound(self) -> None:
        output = BytesIO()
        Image.new("RGB", (4, 4), "white").save(output, format="BMP")
        bmp = output.getvalue()

        self.assertIsNotNone(_convert_raster_to_png(bmp))
        with patch("services.chunker.stages.rasterizer.MAX_RENDER_PIXELS_PER_PAGE", 15):
            self.assertIsNone(_convert_raster_to_png(bmp))


if __name__ == "__main__":
    unittest.main()
