"""The XLSX reader keeps a workbook's layout and decides nothing semantic.

It applies the DOCX table rules - first row of a table supplies the headers, each later
row is one passage - and carries the sheet, region title and cell range on every passage,
so each tool's model reads a row with what it is about and a trace points to its cells.
"""

from __future__ import annotations

import datetime
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image

from services.chunker import run_pipeline
from services.chunker.pipeline import parse_context_file
from services.chunker.stages import parser, parser_xlsx
from shared.document_metadata import extraction_context
from shared.word_export import passage_location


class XlsxReaderTests(unittest.TestCase):
    def setUp(self) -> None:
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.directory = Path(directory.name)

    def read(self, workbook: Workbook, name: str = "plan.xlsx"):
        path = self.directory / name
        workbook.save(path)
        return run_pipeline(str(path), "plan")

    def targets(self) -> Workbook:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Efficacy targets"
        sheet["A1"] = "Target product profile"
        sheet.append(["Attribute", "Minimum", "Preferred", "Notes"])
        sheet.append(["Efficacy", 0.7, 0.9, "against clinical malaria"])
        sheet["B3"].number_format = "0%"
        sheet["C3"].number_format = "0%"
        sheet.append(["Doses", 3, 1, None])
        sheet.append([])
        sheet.append(["Prepared by the team; figures are draft."])
        sheet.append([])
        sheet.append(["Start", datetime.datetime(2026, 3, 1)])
        return workbook

    def test_a_table_reads_as_titled_rows_carrying_their_headers(self) -> None:
        title, efficacy, doses, note, start = self.read(self.targets())

        self.assertEqual((title.block_type, title.content), ("heading", "Target product profile"))
        self.assertEqual(efficacy.block_type, "table_row")
        self.assertEqual(
            efficacy.content,
            "Attribute: Efficacy, Minimum: 70%, Preferred: 90%, Notes: against clinical malaria",
        )
        self.assertEqual(efficacy.heading_stack, ["Efficacy targets", "Target product profile"])
        self.assertEqual(efficacy.structural_meta["cell_range"], "A3:D3")
        self.assertEqual(efficacy.structural_meta["sheet"], "Efficacy targets")
        self.assertEqual(
            [cell["header"] for cell in efficacy.structural_meta["table_cells"]],
            ["Attribute", "Minimum", "Preferred", "Notes"],
        )
        # A blank cell is left out, as a DOCX row's is.
        self.assertEqual(doses.content, "Attribute: Doses, Minimum: 3, Preferred: 1")
        self.assertEqual((note.block_type, note.content), ("paragraph", "Prepared by the team; figures are draft."))
        self.assertEqual(start.content, "Start | 2026-03-01")
        self.assertEqual({block.source_format for block in (title, efficacy, start)}, {"xlsx"})
        self.assertEqual([block.id for block in (title, efficacy)], ["plan/b-0000", "plan/b-0001"])

    def test_a_region_is_layout_so_it_joins_no_blocks(self) -> None:
        """Screener expands a selection along parser-declared tables; blank-row regions
        are not declared, and expanding one could pull a whole raw-data sheet in."""
        for block in self.read(self.targets()):
            self.assertNotIn("table_index", block.structural_meta)
            self.assertNotIn("table_group", block.structural_meta)

    def test_a_merged_title_is_one_cell_and_merged_values_repeat(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.merge_cells("A1:C1")
        sheet["A1"] = "Release specification"
        sheet.append(["Test", "Limit", "Method"])
        sheet.append(["Assay", "95-105%", "HPLC"])
        sheet.append(["Purity", "≥ 98%", None])
        sheet.merge_cells("C3:C4")
        title, assay, purity = self.read(workbook)

        self.assertEqual(title.content, "Release specification")
        self.assertEqual(title.structural_meta["cell_range"], "A1:C1")
        self.assertEqual(purity.content, "Test: Purity, Limit: ≥ 98%, Method: HPLC")

    def test_a_group_label_over_two_header_rows_joins_their_labels(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet["C1"] = "Targets"
        sheet.merge_cells("C1:D1")
        for column, text in zip("ABCD", ["Category", "Attribute", "Minimum", "Preferred"]):
            sheet[f"{column}2"] = text
        sheet.append(["Clinical", "Efficacy", "50%", "75%"])
        (row,) = self.read(workbook)
        self.assertEqual(
            row.content, "Category: Clinical, Attribute: Efficacy, Targets › Minimum: 50%, Targets › Preferred: 75%",
        )
        self.assertEqual(row.structural_meta["header_range"], "A1:D2")

    def test_a_side_table_is_its_own_table_but_a_spacer_column_splits_nothing(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        for row in (["Attribute", None, "Minimum", None, "Assumption", "Value"],
                    ["Efficacy", None, "50%", None, "Coverage", "80%"],
                    ["Doses", None, "3", None, None, None]):
            sheet.append(row)
        efficacy, doses, coverage = self.read(workbook)
        # Column B is a spacer: both sides cover rows 1-3, so they stay one table.
        self.assertEqual(efficacy.content, "Attribute: Efficacy, Minimum: 50%")
        self.assertEqual(doses.structural_meta["cell_range"], "A3:C3")
        # Columns E-F cover rows 1-2 only: a table of its own, not part of row 2.
        self.assertEqual((coverage.content, coverage.structural_meta["cell_range"]), ("Assumption: Coverage, Value: 80%", "E2:F2"))

    def test_a_hidden_row_is_read_and_marked(self) -> None:
        workbook = self.targets()
        workbook.active.row_dimensions[4].hidden = True
        doses = next(block for block in self.read(workbook) if block.content.startswith("Attribute: Doses"))
        self.assertIs(doses.structural_meta["hidden_row"], True)

    def test_one_column_and_one_row_regions_are_paragraphs(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["First note"])
        sheet.append(["Second note"])
        first, second = self.read(workbook)
        self.assertEqual([first.content, second.content], ["First note", "Second note"])
        self.assertEqual(second.structural_meta["cell_range"], "A2")

    def test_a_hidden_sheet_is_read_and_marked(self) -> None:
        workbook = self.targets()
        hidden = workbook.create_sheet("Calculations")
        hidden.sheet_state = "hidden"
        hidden.append(["Total", 42])
        block = self.read(workbook)[-1]
        self.assertEqual(block.heading_stack, ["Calculations"])
        self.assertIs(block.structural_meta["hidden_sheet"], True)

    def test_a_workbook_whose_formulas_were_never_calculated_is_refused(self) -> None:
        """Reading it would report every computed value as empty."""
        workbook = Workbook()
        workbook.active.append([1, 2, "=A1+B1"])
        with self.assertRaisesRegex(ValueError, "plan: the workbook's formulas have no saved results"):
            self.read(workbook)

    def test_an_empty_or_oversized_workbook_is_refused_whole(self) -> None:
        with self.assertRaisesRegex(ValueError, "plan: workbook has no filled cells"):
            self.read(Workbook())
        with patch.object(parser_xlsx, "MAX_XLSX_ROWS", 3), self.assertRaisesRegex(
            ValueError, "more than 3 filled rows"
        ):
            self.read(self.targets())

    def test_a_sheet_stretched_by_formatting_is_refused_before_its_cells_are_walked(self) -> None:
        workbook = self.targets()
        workbook.active["Z5000"].number_format = "0%"
        with patch.object(parser_xlsx, "MAX_SHEET_CELLS", 1_000), patch.object(
            parser_xlsx, "_require_calculated", side_effect=AssertionError("walked the cells first")
        ), self.assertRaisesRegex(ValueError, "spans too many cells"):
            self.read(workbook)

    def test_an_unreadable_file_names_the_document_not_the_temporary_path(self) -> None:
        path = self.directory / "upload-temp.xlsx"
        path.write_bytes(b"not a workbook")
        with self.assertRaises(ValueError) as caught:
            run_pipeline(str(path), "plan")
        self.assertIn("plan: not a readable XLSX workbook", str(caught.exception))
        self.assertNotIn(str(path), str(caught.exception))

    def test_images_are_retained_and_charts_are_named_as_unsupported(self) -> None:
        workbook = self.targets()
        sheet = workbook.active
        buffer = io.BytesIO()
        Image.new("RGB", (8, 8), "red").save(buffer, format="PNG")
        picture = SheetImage(io.BytesIO(buffer.getvalue()))
        sheet.add_image(picture, "F2")
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=2, min_row=2, max_row=4), titles_from_data=True)
        sheet.add_chart(chart, "F10")
        blocks = self.read(workbook)
        image = next(block for block in blocks if block.image is not None)
        self.assertEqual((image.block_type, image.content), ("image", "[image]"))
        self.assertEqual(image.structural_meta["cell_range"], "F2")
        chart_block = next(
            block for block in blocks
            if "unsupported_document_visual" in block.structural_meta.get("extraction_warnings", [])
        )
        self.assertEqual(chart_block.structural_meta["unsupported_visual_kind"], "chart")

    def test_a_legacy_workbook_is_converted_then_read_by_the_same_reader(self) -> None:
        converted = self.directory / "converted.xlsx"
        self.targets().save(converted)
        legacy = self.directory / "plan.xls"
        legacy.write_bytes(b"legacy")
        with patch.object(parser, "convert_legacy_document", return_value=str(converted)) as convert:
            blocks = run_pipeline(str(legacy), "plan")
        convert.assert_called_once()
        self.assertEqual({block.source_format for block in blocks}, {"xls"})
        self.assertEqual(blocks[1].structural_meta["cell_range"], "A3:D3")

    def test_ask_reads_a_workbook_with_the_same_reader(self) -> None:
        path = self.directory / "plan.xlsx"
        self.targets().save(path)
        self.assertEqual(
            [block.content for block in parse_context_file(str(path), "plan")],
            [block.content for block in run_pipeline(str(path), "plan")],
        )


class WorkbookLocationTests(unittest.TestCase):
    """Where a workbook passage sits, said the same way to a model, an export and Ask."""

    META = {"sheet": "Targets", "sheet_index": 2, "cell_range": "B4:F4", "hidden_sheet": True, "hidden_row": True}

    def test_a_model_is_told_the_sheet_and_cells(self) -> None:
        self.assertEqual(extraction_context(self.META), "sheet=Targets; hidden_sheet=true; hidden_row=true; cells=B4:F4")

    def test_an_export_names_it_as_excel_writes_a_reference(self) -> None:
        self.assertEqual(passage_location({"structural_meta": self.META, "heading_stack": ["Targets"]}), "Targets!B4:F4")

    def test_ask_places_a_visual_by_sheet(self) -> None:
        from services.assistant.workspace import location_of

        self.assertEqual(location_of({"structural_meta": self.META}), "sheet 2")


if __name__ == "__main__":
    unittest.main()
