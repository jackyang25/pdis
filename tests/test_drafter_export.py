"""The Word table a team takes away: each suggested target and the sources behind it."""

from __future__ import annotations

import io
import unittest

from docx import Document

from services.drafter import export_report
from shared.word_export import word_document
from tests.test_drafter_contract import result, suggested


def export_docx(draft: dict) -> bytes:
    return word_document(export_report(draft))


def table_rows(data: bytes) -> list[list[str]]:
    document = Document(io.BytesIO(data))
    return [[cell.text for cell in row.cells] for row in document.tables[0].rows]


def insufficient(name: str) -> dict:
    return suggested(name=name, suggestion={
        "kind": "insufficient", "minimum": "", "preferred": "", "statement": "",
        "rationale": "", "leans_on": "", "basis_ids": [],
    }, observations=[])


class ExportTests(unittest.TestCase):
    def test_each_row_shows_the_suggestion_what_it_leans_on_and_its_sources(self):
        failed = suggested(name="vaccine.shelf_life", status="failed", suggestion=None,
                           failure={"code": "suggestion_failed", "message": "No valid reply.", "cause": "invalid_model_response"})
        rows = table_rows(export_docx(result(suggested(), insufficient("vaccine.cogs"), failed)))
        self.assertEqual(rows[0], ["Variable", "Suggested target", "Based on", "Sources"])
        self.assertEqual(rows[1][1], "Minimum: 50%\nPreferred: 75%")
        self.assertEqual(rows[1][2], "Comparable products and published guidance")
        self.assertIn("Trial A (PubMed, 2021)", rows[1][3])
        self.assertEqual(rows[2][1:3], ["Not set", "Not enough evidence"])
        self.assertEqual(rows[3][1:3], ["Not set", "Could not be suggested"])

    def test_the_file_shows_what_steered_the_draft(self):
        document = Document(io.BytesIO(export_docx({**result(), "intended_use": "Maternal immunization to protect infants in their first six months."})))
        self.assertIn("Intended use, as the team stated it: Maternal immunization to protect infants in their first six months.", [paragraph.text for paragraph in document.paragraphs])

    def test_the_document_says_its_targets_are_suggestions(self):
        document = Document(io.BytesIO(export_docx(result())))
        self.assertTrue(any("Every target is a suggestion" in paragraph.text for paragraph in document.paragraphs))


if __name__ == "__main__":
    unittest.main()
