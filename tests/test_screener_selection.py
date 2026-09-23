"""Source selection preserves evidence rather than generating replacement prose."""
from __future__ import annotations

import unittest
import tempfile
import zipfile
from pathlib import Path

from services.chunker import ContentBlock, ImageAsset
from services.screener import QuestionSpec
from shared.errors import ModelResponseError


def block(identifier, content="source", *, doc="report", meta=None, image=None):
    return ContentBlock(identifier, doc, 0, "image" if image else "paragraph",
                        content, [], meta or {}, {}, image=image)


class SelectionClient:
    def __init__(self, *replies):
        self.replies = list(replies)
        self.calls = []

    def call_structured(self, system_prompt, user_message, max_tokens, **kwargs):
        self.calls.append((system_prompt, user_message, kwargs))
        reply = self.replies.pop(0)
        if isinstance(reply, Exception):
            raise reply
        return reply


class SelectionTests(unittest.TestCase):
    def test_distinct_textbox_tables_in_same_part_stay_separate(self):
        from docx import Document
        from docx.oxml import parse_xml
        from services.chunker import run_pipeline
        from services.screener.evidence import expand_selection
        from tests.test_chunker_docx_parts import W, V

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "textboxes.docx"
            document = Document()
            cell = document.add_table(rows=1, cols=1).cell(0, 0)
            for name in ("Evidence A", "Unrelated B"):
                paragraph = cell.add_paragraph("Anchor")
                paragraph._p.append(parse_xml(
                    f'<w:r xmlns:w="{W}" xmlns:v="{V}"><w:pict><v:shape><v:textbox>'
                    '<w:txbxContent><w:tbl><w:tblPr/><w:tblGrid><w:gridCol/></w:tblGrid>'
                    f'<w:tr><w:tc><w:tcPr/><w:p><w:r><w:t>{name}</w:t></w:r></w:p>'
                    '</w:tc></w:tr></w:tbl></w:txbxContent></v:textbox></v:shape></w:pict></w:r>'
                ))
            document.add_table(rows=1, cols=1).cell(0, 0).text = "Unrelated body table"
            document.save(path)
            blocks = run_pipeline(str(path), doc_id="tables", org="bmgf",
                                  intervention_class="drug", indication="hiv")
        evidence = next(b for b in blocks if b.content == "Evidence A")
        self.assertEqual([b.content for b in expand_selection(blocks, [evidence.id])], ["Evidence A"])

    def test_docx_supplement_table_does_not_expand_into_unrelated_body_table(self):
        from services.chunker import run_pipeline
        from services.screener.evidence import expand_selection
        from tests.test_chunker_docx_parts import _coverage_fixture, _replace_zip_members

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "tables.docx"
            _coverage_fixture(path)
            with zipfile.ZipFile(path) as archive:
                body = archive.read("word/document.xml")
            _replace_zip_members(path, {
                "word/document.xml": body.replace(b'footnoteReference w:id="9"', b'footnoteReference w:id="2"', 1),
            })
            blocks = run_pipeline(str(path), doc_id="tables", org="bmgf",
                                  intervention_class="drug", indication="hiv")
        footnote = next(b for b in blocks if b.content == "Footnote table cell")
        selected = expand_selection(blocks, [footnote.id])
        self.assertNotIn("Before text box", [b.content for b in selected])
        self.assertIn(footnote, selected)
        outer = next(b for b in blocks if b.content == "Outer cell")
        self.assertIn("Nested cell", [b.content for b in expand_selection(blocks, [outer.id])])

    def select(self, client, blocks):
        from services.screener.stages.selector import select_evidence
        return select_evidence(
            QuestionSpec("Q1", "What is the study plan and timeline?", requirement="anticipatory"),
            blocks, indication="tuberculosis", intervention_class="monoclonal_antibody",
            llm_client=client, max_tokens=16000,
        )

    def test_selection_returns_original_blocks_in_source_order(self):
        plan, unrelated, timeline = [block(name) for name in ("plan", "other", "timeline")]
        selected = self.select(SelectionClient({"block_ids": ["timeline", "plan", "plan"]}),
                               [plan, unrelated, timeline])
        self.assertEqual(selected, [plan, timeline])
        self.assertIs(selected[0], plan)
        self.assertIs(selected[1], timeline)

    def test_explicit_groups_keep_table_and_visual_context_without_crossing_documents(self):
        from services.screener.evidence import expand_selection
        for key, value in (("page", 1), ("slide", 2), ("table_index", 0)):
            with self.subTest(key=key):
                first = block("first", meta={key: value})
                second = block("second", meta={key: value})
                unrelated = block("other", meta={key: value + 1})
                other_document = block("foreign", doc="foreign", meta={key: value})
                self.assertEqual(expand_selection([first, second, unrelated, other_document], ["second"]),
                                 [first, second])

    def test_group_expansion_follows_linked_groups_but_not_missing_metadata(self):
        from services.screener.evidence import expand_selection
        blocks = [block("table", meta={"table_index": 0}),
                  block("link", meta={"table_index": 0, "slide": 1}),
                  block("visual", meta={"slide": 1}), block("unrelated")]
        self.assertEqual([b.id for b in expand_selection(blocks, ["table"])],
                         ["table", "link", "visual"])
        self.assertEqual(expand_selection(blocks, []), [])

    def test_request_reads_whole_document_and_preserves_image_and_warning(self):
        asset = ImageAsset(media_type="image/png", data_base64="aW1hZ2U=",
                           sha256="fixture", source_media_type="image/png")
        blocks = [block("text", "Study context", meta={"page": 1, "extraction_warnings": ["pdf_text_layout"]}),
                  block("image", meta={"page": 1}, image=asset), block("other", "Unrelated content")]
        client = SelectionClient({"block_ids": ["image"]})
        selected = self.select(client, blocks)
        system, message, kwargs = client.calls[0]
        self.assertIn("Study context", message)
        self.assertIn("Unrelated content", message)
        self.assertIn("pdf_text_layout", message)
        self.assertIn("What is the study plan and timeline?", message)
        self.assertIn("monoclonal antibody", message)
        self.assertNotIn("anticipatory", system + message)
        self.assertEqual(kwargs["images"], [{"block_id": "image", "data_url": "data:image/png;base64,aW1hZ2U="}])
        self.assertEqual([b.id for b in selected], ["text", "image"])
        self.assertIs(selected[1].image, asset)

    def test_empty_selection_is_success_not_a_retry(self):
        client = SelectionClient({"block_ids": []})
        self.assertEqual(self.select(client, [block("a")]), [])
        self.assertEqual(len(client.calls), 1)

    def test_invalid_reply_is_repaired_once_then_fails_not_empty_evidence(self):
        for reply in (None, {}, {"block_ids": "a"}, {"block_ids": [1]},
                      {"block_ids": ["foreign:1"]}, {"block_ids": [" a "]},
                      {"block_ids": [], "decision": "answered"}):
            with self.subTest(reply=reply):
                client = SelectionClient(reply, reply)
                with self.assertRaises(ModelResponseError):
                    self.select(client, [block("a")])
                self.assertEqual(len(client.calls), 2)

    def test_contract_repair_can_return_valid_source(self):
        client = SelectionClient({"block_ids": ["foreign"]}, {"block_ids": ["a"]})
        source = block("a")
        self.assertEqual(self.select(client, [source]), [source])
        self.assertIn("prior", client.calls[1][1].lower())

    def test_provider_failure_propagates_without_becoming_selection(self):
        client = SelectionClient(RuntimeError("gateway unavailable"))
        with self.assertRaisesRegex(RuntimeError, "gateway unavailable"):
            self.select(client, [block("a")])
        self.assertEqual(len(client.calls), 1)

    def test_selector_refuses_empty_or_multiple_documents_before_model(self):
        for blocks in ([], [block("a"), block("b", doc="another")]):
            client = SelectionClient()
            with self.assertRaises(ValueError):
                self.select(client, blocks)
            self.assertEqual(client.calls, [])
