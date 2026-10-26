"""Every tool's Word export: one layout, each tool's rows in its page's order and words."""

from __future__ import annotations

import io
import re
import unittest
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from api import schemas
from services import aligner, inspector, scout, screener
from shared.word_export import (
    Table,
    WordReport,
    count_summary,
    passage_location,
    passage_sources,
    word_document,
)

WEB = Path(__file__).resolve().parents[1] / "web"


def read(data: bytes) -> tuple[list[str], list[list[list[str]]]]:
    document = Document(io.BytesIO(data))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    return paragraphs, tables


def block(block_id: str, doc_id: str = "plan", **meta) -> dict:
    section = meta.pop("section", None)
    return {"id": block_id, "doc_id": doc_id, "ordinal": 0, "block_type": "paragraph", "content": "x",
            "heading_stack": [], "section_label": section, "structural_meta": meta, "style_hint": {}}


class LayoutTests(unittest.TestCase):
    def test_every_export_shares_one_page_and_one_caveat(self) -> None:
        data = word_document(WordReport(
            title="T", context="C", about="A", summary="S",
            tables=[Table(heading="Empty", note="Nothing here.", columns=("A",), rows=[]),
                    Table(columns=("A", "B"), rows=[("1", "")], widths=(1, 3))],
        ))
        document = Document(io.BytesIO(data))
        self.assertGreater(document.sections[0].page_width, document.sections[0].page_height)
        narrow, wide = document.tables[0].columns
        self.assertAlmostEqual(wide.width / narrow.width, 3, places=2)
        paragraphs, tables = read(data)
        self.assertIn("Verify findings against the cited sources", " ".join(paragraphs))
        # A table with no rows is its note alone, and an empty cell says it is empty.
        self.assertEqual(tables, [[["A", "B"], ["1", "—"]]])
        self.assertIn("Nothing here.", paragraphs)

    def test_counts_read_as_a_sentence(self) -> None:
        self.assertEqual(count_summary(2, ("item", "items"), [(1, "Specified"), (1, "N/A")]),
                         "2 items: 1 specified, 1 N/A.")

    def test_a_passage_is_named_where_a_reader_finds_it(self) -> None:
        self.assertEqual(passage_location(block("b", page=4)), "Page 4")
        self.assertEqual(passage_location(block("b", slide=3)), "Slide 3")
        self.assertEqual(passage_location(block("b", section="Executive Summary")), "Executive Summary")
        self.assertEqual(passage_location(block("b", document_part_kind="footnote", note_id=2)), "Footnote 2")
        blocks = {"a": block("a", "the_plan", page=1), "b": block("b", "the_plan", page=1)}
        # One line per place, its document named when the result reads several, and a
        # citation the file no longer holds said to be unavailable rather than dropped.
        self.assertEqual(passage_sources(["a", "b", "gone"], blocks, name_documents=True),
                         "the plan · Page 1\nPassage not in the saved result")


class ScreenerExportTests(unittest.TestCase):
    def review(self) -> dict:
        question = lambda qid, state, requirement="required", **extra: {
            "id": qid, "text": f"Question {qid}?", "state": state, "requirement": requirement,
            "statement": extra.get("statement", "Covered." if state != "not_found" else ""),
            "missing": extra.get("missing", ""), "cited_block_ids": ["b1"] if state != "not_found" else [],
        }
        return schemas.GateReviewOut.model_validate({
            "gate_id": "lcs", "gate_label": "Lead Chemical Series Selection", "bank_source": "Bank v2",
            "documents": [{"doc_id": "plan"}], "org": "bmgf", "intervention_class": "drug",
            "indication": "female_contraception", "blocks": [block("b1", page=2)],
            "disciplines": [
                {"id": "cmc", "label": "CMC", "questions": [
                    question("CMC.1", "partly_answered", "anticipatory", missing="Impurities."),
                    question("CMC.2", "partly_answered", missing="Stability."),
                    question("CMC.3", "not_applicable"),
                ]},
                {"id": "reg", "label": "Regulatory", "questions": [question("REG.1", "not_found")]},
            ],
        }).model_dump()

    def test_the_page_panels_in_order_with_required_questions_first(self) -> None:
        paragraphs, tables = read(screener.export_docx(self.review()))
        self.assertIn("3 questions: 0 answered, 2 partly answered, 1 not found.", paragraphs)
        self.assertIn("2 questions this gate requires are still unanswered.", paragraphs)
        self.assertIn("1 question does not apply to this product and is not listed.", paragraphs)
        partly, not_found = tables
        self.assertEqual([row[1] for row in partly[1:]], ["Question CMC.2? (CMC.2)", "Question CMC.1? (CMC.1)"])
        self.assertEqual(partly[1][3], "Covered.\nStill not stated: Stability.")
        self.assertEqual(partly[1][4], "Page 2")
        self.assertEqual(not_found[1][:4], ["Regulatory", "Question REG.1? (REG.1)", "Required", "—"])


class AlignerExportTests(unittest.TestCase):
    def test_each_comparison_lists_shortfalls_first_with_both_documents_cited(self) -> None:
        finding = lambda rid, verdict, statement="It says so.": {
            "requirement_id": rid, "edge_id": "itpp-to-ctpp", "requirement": f"Requirement {rid}.",
            "reference_spans": [{"quote": "q", "block_ids": ["r1"]}], "verdict": verdict,
            "statement": statement, "comparison_spans": [{"quote": "q", "block_ids": ["c1"]}] if statement else [],
        }
        alignment = schemas.AlignmentResultOut.model_validate({
            "documents": [{"doc_id": "profile", "source_type": "itpp", "display_name": "Intervention TPP"},
                          {"doc_id": "candidate", "source_type": "ctpp", "display_name": "Candidate TPP"}],
            "edges": [{"edge_id": "itpp-to-ctpp", "reference_doc_id": "profile",
                       "comparison_doc_id": "candidate", "question": "Does it meet the bar?"}],
            "org": "bmgf", "intervention_class": "vaccine", "indication": "malaria",
            "blocks": [block("r1", "profile", section="Efficacy"), block("c1", "candidate", page=3)],
            "findings": [finding("r1", "meets"), finding("r2", "not_addressed", ""), finding("r3", "falls_short")],
        }).model_dump()
        paragraphs, (table,) = read(aligner.export_docx(alignment))
        self.assertIn("Intervention TPP → Candidate TPP", paragraphs)
        self.assertEqual(table[0], ["Requirement", "Verdict", "What the cTPP says", "Sources"])
        self.assertEqual([row[1] for row in table[1:]], ["Falls short", "Not addressed", "Meets"])
        self.assertEqual(table[1][3], "profile · Efficacy\ncandidate · Page 3")
        self.assertEqual(table[2][2], "—")


class InspectorExportTests(unittest.TestCase):
    def inspection(self, **overrides) -> dict:
        unit = lambda uid, verdict, variable=None, statement="": {
            "id": uid, "verdict": verdict, "statement": statement, "section_name": "Profile",
            "variable_name": variable, "cited_block_ids": ["b1"] if verdict not in {"not_present", "not_applicable"} else [],
        }
        payload = {
            "doc_id": "rsv_itpp", "applicability_facts": {}, "consistency_status": "complete",
            "consistency_failure_code": None, "assessment_status": "complete", "org": "bmgf",
            "source_type": "itpp", "intervention_class": "vaccine", "indication": "malaria",
            "blocks": [block("b1", "rsv_itpp", section="Profile")], "document_findings": [],
            "rubric_resolutions": [
                {"rubric_id": "pdid", "display_name": "PDID template", "status": "included", "reason_code": "", "reason": ""},
                {"rubric_id": "who", "display_name": "WHO presentation", "status": "needs_context",
                 "reason_code": "fact_unknown", "reason": "Whether the vaccine is multi-dose is not known."},
            ],
            "reviews": [{"assessment_status": "complete", "rubric": {
                "id": "pdid", "revision": "1.0", "updated_on": None, "display_name": "PDID template",
                "authority": "PDID", "scope": "", "stage_guidance": "", "mirrors": None,
                "evidence_scope": "mapped_section", "sources": [], "requirements": [],
            }, "sections": [{"section_name": "Profile", "mapped_block_ids": ["b1"], "is_present": True,
                             "verdict_counts": {}, "units": [
                                 unit("u1", "specified"),
                                 unit("u2", "insufficient", "Efficacy", "No endpoint is named."),
                             ]}]}],
            **overrides,
        }
        return schemas.InspectionResultOut.model_validate(payload).model_dump()

    def test_every_requirement_keeps_its_row_and_unrun_rubrics_say_why(self) -> None:
        paragraphs, (review,) = read(inspector.export_docx(self.inspection()))
        self.assertIn("WHO presentation was not run: Whether the vaccine is multi-dose is not known.", paragraphs)
        self.assertIn("Revision 1.0. 2 requirements: 1 specified, 0 not present, 0 placeholder, 1 insufficient, 0 vague.",
                      paragraphs)
        self.assertEqual(review[1][:3], ["Profile", "Specified", "—"])
        self.assertEqual(review[2][:3], ["Profile › Efficacy", "Insufficient", "No endpoint is named."])
        self.assertIn("No conflicts between sections were found.", paragraphs)

    def test_a_failed_consistency_check_is_not_read_as_no_conflict(self) -> None:
        paragraphs, _ = read(inspector.export_docx(self.inspection(
            consistency_status="failed", consistency_failure_code="response_truncated")))
        self.assertTrue(any("did not complete" in paragraph for paragraph in paragraphs))
        self.assertNotIn("No conflicts between sections were found.", paragraphs)


class ScoutExportTests(unittest.TestCase):
    def result(self, phase: str = "final") -> dict:
        finding = {"url": "https://pubmed.test/1", "title": "Trial A", "query": "q",
                   "retrieved_at": "2026-01-01T00:00:00Z", "source": "pubmed",
                   "source_labels": {"pubmed": "PubMed"}, "published_at": "2021-04-01"}
        assessment = lambda ref, strength, reason="", findings=(): {
            "attribute_ref": ref, "strength": strength, "reason": reason, "doc_target": f"{ref} target",
            "supporting_findings": list(findings), "failure_code": "provider_error" if strength == "not_assessed" else "",
        }
        payload = {
            "phase": phase, "org": "bmgf", "source_type": "itpp", "intervention_class": "vaccine",
            "indication": "malaria",
            "context_validation": {"status": "match", "configured_indication": "malaria"},
            "retrieval_scope": {"entries": []},
            "quantitative_ledger": {"status": "complete"}, "matches": [],
            "variables": [{"name": "vaccine.efficacy", "description": "", "document_spans": [
                {"quote": "Variable: Efficacy, Minimum The minimal target: 70%", "block_ids": ["row"]}]}],
            "assessments": [
                assessment("vaccine.efficacy", "well_grounded", "Two trials met it.", [finding]),
                assessment("vaccine.safety", "not_assessed", "Its evidence could not be read: Request timed out."),
                assessment("vaccine.cogs", "not_searched"),
            ],
            "precedents": [{"attribute_ref": "vaccine.efficacy", "precedent": "direct", "outcome": "mixed",
                            "reason": "One approval, one halt."},
                           {"attribute_ref": "vaccine.safety", "precedent": "not_assessed", "outcome": "not_assessed",
                            "failure_code": "provider_error"}],
            "conformity": [{"attribute_refs": ["vaccine.efficacy"], "target_id": "t1", "target_role": "threshold",
                            "target_value": 80.0, "comparator": ">=", "unit": "%", "target_label": "Efficacy",
                            "target_meeting_count": 2, "target_meeting_rate": 0.66, "verdict": "",
                            "benchmark_count": 3, "benchmark_minimum": 62.5, "benchmark_maximum": 91.0,
                            "benchmark_median": 84.0, "calibration_status": "limited"}],
            "stats": {"queries": 1, "findings": 1, "unique_findings": 1, "insights": 1, "matches": 0, "assessments": 3},
            "blocks": [block("b1", "rsv_itpp"), {
                **block("row", "rsv_itpp"), "block_type": "table_row",
                "content": "Variable: Efficacy, Minimum\nThe minimal target: 70%",
                "structural_meta": {"column_headers": ["Variable", "Minimum\nThe minimal target"], "table_cells": [
                    {"column_index": 0, "header": "Variable", "value": "Efficacy", "content_start": 0,
                     "content_end": 18, "value_start": 10, "value_end": 18},
                    {"column_index": 1, "header": "Minimum\nThe minimal target", "value": "70%", "content_start": 20,
                     "content_end": 51, "value_start": 48, "value_end": 51}]},
            }],
        }
        return schemas.ScoutRunResponse.model_validate(payload).model_dump()

    def test_each_field_keeps_its_verdicts_reasons_and_sources(self) -> None:
        paragraphs, (fields, numbers) = read(scout.export_docx(self.result()))
        self.assertIn("Evidence review: rsv itpp", paragraphs)
        efficacy, safety, cogs = fields[1:]
        # A table row by its cells, each under its heading's own title line.
        self.assertEqual(efficacy[:4], ["Efficacy", "Variable: Efficacy\nMinimum: 70%", "Well grounded\nTwo trials met it.",
                                        "Direct precedent · Mixed outcome\nOne approval, one halt."])
        self.assertEqual(efficacy[4], "Trial A (PubMed, 2021) https://pubmed.test/1")
        # A failure reads as one, never as the pipeline's diagnostic or as a finding.
        self.assertTrue(safety[2].startswith("Grounding not assessed\nNo verdict was obtained"))
        self.assertTrue(safety[3].startswith("Precedent not assessed\n"))
        self.assertTrue(cogs[2].startswith("Not searched\nExcluded at review"))
        self.assertEqual(numbers[1], ["Efficacy: >= 80 %", "3", "62.5–91 %", "84 %", "Limited"])
        self.assertTrue(any("not confidence intervals" in paragraph for paragraph in paragraphs))

    def test_a_review_still_being_decided_does_not_export(self) -> None:
        with self.assertRaisesRegex(ValueError, "final"):
            scout.export_docx(self.result(phase="evidence_review"))


class RouteTests(unittest.TestCase):
    def test_each_tool_exports_a_word_file_and_refuses_an_unfinished_review(self) -> None:
        from api.main import app

        client = TestClient(app)
        response = client.post("/api/screener/export", json={"result": ScreenerExportTests().review()})
        self.assertEqual(response.status_code, 200)
        self.assertIn('filename="pdis-screener.docx"', response.headers["content-disposition"])
        self.assertTrue(response.content.startswith(b"PK"))
        refused = client.post("/api/scout/export", json={"result": ScoutExportTests().result("evidence_review")})
        self.assertEqual(refused.status_code, 422)


class LabelParityTests(unittest.TestCase):
    """A Word file names every state as the page does: the labels are copied, so they are compared."""

    def ts_record(self, path: str, name: str) -> dict[str, str]:
        source = (WEB / path).read_text(encoding="utf-8")
        body = re.search(rf"export const {name}\b[^=]*=\s*\{{(.*?)\n\}};", source, re.S)
        self.assertIsNotNone(body, f"{name} not found in {path}")
        return dict(re.findall(r'^\s*(\w+): "([^"]+)",', body.group(1), re.M))

    def test_labels_match_the_web(self) -> None:
        pairs = [
            (inspector.export.VERDICT_LABEL, "lib/api.ts", "VERDICT_LABEL"),
            (aligner.export.VERDICT_LABEL, "lib/api.ts", "VERDICT_LABELS"),
            (screener.export.STATE_LABEL, "lib/api.ts", "QUESTION_STATE_LABEL"),
            (screener.export.REQUIREMENT_LABEL, "lib/api.ts", "QUESTION_REQUIREMENT_LABEL"),
            (scout.export.GROUNDING_LABEL, "lib/scout-labels.ts", "GROUNDING_LABEL"),
            (scout.export.PRECEDENT_LABEL, "lib/scout-labels.ts", "PRECEDENT_LABEL"),
            (scout.export.CALIBRATION_LABEL, "lib/scout-labels.ts", "CALIBRATION_BASIS_LABEL"),
        ]
        for labels, path, name in pairs:
            with self.subTest(name=name):
                self.assertEqual(labels, self.ts_record(path, name))
        web_outcomes = self.ts_record("lib/scout-labels.ts", "OUTCOME_LABEL")
        self.assertEqual(scout.export.OUTCOME_LABEL, {key: web_outcomes[key] for key in scout.export.OUTCOME_LABEL})


if __name__ == "__main__":
    unittest.main()
