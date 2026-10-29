"""Every tool's Excel export: the same report as its Word file, laid out for sorting and filtering."""

from __future__ import annotations

import io
import unittest

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from services import aligner, drafter, inspector, scout, screener
from shared.report_export import ExportReport, Table
from shared.word_export import word_document
from shared.xlsx_export import MAX_SHEET_NAME, xlsx_workbook
from tests.test_drafter_contract import result as draft_result, suggested
# The module, not its classes: a TestCase imported by name is collected and run again here.
from tests import test_word_exports as word_exports


def sheets(data: bytes) -> dict[str, list[list[str]]]:
    workbook = load_workbook(io.BytesIO(data))
    return {sheet.title: [["" if value is None else str(value) for value in row] for row in sheet.iter_rows(values_only=True)]
            for sheet in workbook.worksheets}


class LayoutTests(unittest.TestCase):
    def report(self) -> ExportReport:
        return ExportReport(
            title="T", context="C", about="A", summary="S", notes=["N"],
            tables=[Table(heading="Empty", note="Nothing here.", columns=("A",), rows=[]),
                    Table(heading="Findings", note="2 findings.", columns=("A", "B"), rows=[("1", ""), ("2", "x\ny")],
                          widths=(1, 3))],
        )

    def test_the_summary_sheet_opens_with_what_the_word_file_opens_with(self) -> None:
        summary = sheets(xlsx_workbook(self.report()))["Summary"]
        lines = [row[0] for row in summary]
        self.assertEqual(lines[:3], ["T", "C", "A"])
        self.assertIn("Verify findings against the cited sources", lines[3])
        self.assertEqual(lines[4:], ["S", "N"])

    def test_each_table_is_a_sheet_with_a_header_that_stays_in_view_and_filters(self) -> None:
        workbook = load_workbook(io.BytesIO(xlsx_workbook(self.report())))
        self.assertEqual(workbook.sheetnames, ["Summary", "Empty", "Findings"])
        findings = workbook["Findings"]
        rows = [[value for value in row] for row in findings.iter_rows(values_only=True)]
        # The heading, the note, then the header and rows; an empty cell says it is empty, as
        # in Word, and a value keeps its line breaks.
        self.assertEqual(rows, [["Findings", None], ["2 findings.", None], ["A", "B"], ["1", "—"], ["2", "x\ny"]])
        self.assertTrue(findings["A3"].font.bold)
        self.assertEqual(findings.freeze_panes, "A4")
        self.assertEqual(findings.auto_filter.ref, "A3:B5")
        self.assertTrue(findings["B5"].alignment.wrap_text)
        # Printed like the Word file: landscape, one page wide, the header repeated.
        self.assertEqual(findings.page_setup.orientation, "landscape")
        self.assertTrue(findings.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(findings.print_title_rows, "$3:$3")
        self.assertAlmostEqual(findings.column_dimensions["B"].width / findings.column_dimensions["A"].width, 3, delta=0.1)
        # A table with no rows is its note alone: a header over nothing reads as lost rows.
        self.assertEqual([[value for value in row] for row in workbook["Empty"].iter_rows(values_only=True)],
                         [["Empty"], ["Nothing here."]])

    def test_a_sheet_name_keeps_the_heading_within_excel_s_limits(self) -> None:
        long = "Intervention TPP → Candidate TPP: does the candidate meet the bar?"
        report = ExportReport(title="T", context="", about="", tables=[
            Table(heading=long, columns=("A",), rows=[("1",)]),
            Table(heading=long, columns=("A",), rows=[("2",)]),
            Table(columns=("A",), rows=[("3",)]),
        ])
        workbook = load_workbook(io.BytesIO(xlsx_workbook(report)))
        names = workbook.sheetnames[1:]
        self.assertEqual(len(set(names)), 3)
        self.assertTrue(all(len(name) <= MAX_SHEET_NAME and ":" not in name and "?" not in name for name in names))
        # Cut at a word, never mid-word, and the sheet still opens with the heading whole.
        self.assertEqual(names[0], "Intervention TPP → Candidate…")
        self.assertTrue(names[1].endswith("(2)"))
        self.assertEqual(names[2], "Table 3")
        self.assertEqual(workbook[names[0]]["A1"].value, long)

    def test_a_report_s_only_table_is_named_for_the_report(self) -> None:
        report = ExportReport(title="Draft iTPP targets", context="", about="", tables=[Table(columns=("A",), rows=[("1",)])])
        self.assertEqual(load_workbook(io.BytesIO(xlsx_workbook(report))).sheetnames, ["Summary", "Draft iTPP targets"])


class ParityTests(unittest.TestCase):
    """One report behind both files: every table's header and rows read the same in Excel."""

    def assert_same_tables(self, report: ExportReport) -> None:
        _, word_tables = word_exports.read(word_document(report))
        book = sheets(xlsx_workbook(report))
        # Each sheet opens with its heading and any note; what follows is the Word table.
        excel_tables = [
            rows[2:] if table.note else rows[1:]
            for table, rows in zip(report.tables, list(book.values())[1:], strict=True)
            if table.rows
        ]
        self.assertEqual(excel_tables, word_tables)

    def test_every_tool_s_workbook_holds_its_word_tables(self) -> None:
        for tool, result in (
            (screener, word_exports.ScreenerExportTests().review()),
            (aligner, word_exports.alignment_result()),
            (inspector, word_exports.InspectorExportTests().inspection()),
            (scout, word_exports.ScoutExportTests().result()),
            (drafter, draft_result(suggested())),
        ):
            with self.subTest(tool=tool.__name__):
                self.assert_same_tables(tool.export_report(result))


class RouteTests(unittest.TestCase):
    def test_format_selects_the_file_and_an_unknown_format_is_refused(self) -> None:
        from api.main import app

        client = TestClient(app)
        review = {"result": word_exports.ScreenerExportTests().review()}
        workbook = client.post("/api/screener/export?format=xlsx", json=review)
        self.assertEqual(workbook.status_code, 200)
        self.assertIn("spreadsheetml", workbook.headers["content-type"])
        self.assertIn('filename="pdis-screener.xlsx"', workbook.headers["content-disposition"])
        load_workbook(io.BytesIO(workbook.content))
        # Word stays the default, so an existing caller keeps the file it asked for.
        self.assertIn("wordprocessingml", client.post("/api/screener/export", json=review).headers["content-type"])
        self.assertEqual(client.post("/api/screener/export?format=csv", json=review).status_code, 422)


if __name__ == "__main__":
    unittest.main()
